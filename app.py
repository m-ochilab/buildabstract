import hashlib
import json
import logging
import os
import queue
import re
import sqlite3
import subprocess
import time
import traceback
from contextlib import closing
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

import httpx
from docx import Document
from dotenv import load_dotenv
from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, PlainTextResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel, Field
from pypdf import PdfReader

logging.basicConfig(level=logging.INFO)


def _load_dotenv(path: Path) -> None:
    if not path.exists():
        return
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, value = line.partition("=")
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


BASE_DIR = Path(__file__).parent
_load_dotenv(BASE_DIR / ".env")
load_dotenv(BASE_DIR / ".env")

DB_PATH = BASE_DIR / "data.db"
GUIDANCE_PATH = Path(os.getenv("GUIDANCE_PATH", BASE_DIR / "guidance.txt"))
MAX_SOURCE_TEXT_LENGTH = 12000
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}
PASS_THRESHOLD = 7.5

# job_id -> accumulated streaming text
_llm_stream: dict[int, str] = {}


def getenv_int(name: str, default: int) -> int:
    raw = os.getenv(name)
    if raw is None:
        return default
    try:
        value = int(raw)
    except ValueError:
        return default
    return value if value > 0 else default


_raw_urls = os.getenv("OLLAMA_URLS", os.getenv("OLLAMA_URL", "http://127.0.0.1:11434"))
OLLAMA_URLS: list[str] = [item.strip() for item in _raw_urls.split(",") if item.strip()]
OLLAMA_URL = OLLAMA_URLS[0]
_ollama_pool: queue.Queue[str] = queue.Queue()
for endpoint in OLLAMA_URLS:
    _ollama_pool.put(endpoint)

DEFAULT_MODEL = os.getenv("OLLAMA_MODEL", "glm-4.7-flash:latest")
REQUEST_TIMEOUT = float(os.getenv("OLLAMA_TIMEOUT_SECONDS", "120"))
SYNOLOGY_CHAT_WEBHOOK_URL = os.getenv("SYNOLOGY_CHAT_WEBHOOK_URL", "").strip()

_raw_or_keys = os.getenv("OPENROUTER_API_KEYS", os.getenv("OPENROUTER_API_KEY", ""))
OPENROUTER_API_KEYS: list[str] = [item.strip() for item in _raw_or_keys.split(",") if item.strip()]
OPENROUTER_MODEL = os.getenv("OPENROUTER_MODEL", "google/gemini-2.5-pro")
OPENROUTER_MAX_CONCURRENCY = getenv_int("OPENROUTER_MAX_CONCURRENCY", 10)
_OPENROUTER_BASE = "https://openrouter.ai/api/v1"
USE_OPENROUTER = bool(OPENROUTER_API_KEYS)
_openrouter_pool: queue.Queue[str] = queue.Queue()
if OPENROUTER_API_KEYS:
    for index in range(OPENROUTER_MAX_CONCURRENCY):
        _openrouter_pool.put(OPENROUTER_API_KEYS[index % len(OPENROUTER_API_KEYS)])


def current_backend() -> str:
    return "openrouter" if USE_OPENROUTER else "ollama"


def model_for_backend(backend_name: str) -> str:
    return OPENROUTER_MODEL if backend_name == "openrouter" else DEFAULT_MODEL


def synology_chat_configured() -> bool:
    return bool(SYNOLOGY_CHAT_WEBHOOK_URL)


def resolve_backend_choice(requested_backend: str) -> str:
    backend_name = requested_backend.strip().lower() or current_backend()
    if backend_name not in {"openrouter", "ollama"}:
        raise HTTPException(status_code=400, detail="不正なバックエンド指定です。")
    if backend_name == "openrouter" and not OPENROUTER_API_KEYS:
        raise HTTPException(status_code=400, detail="OpenRouter の API キーが設定されていません。")
    return backend_name


app = FastAPI(title="BuildAbstract | 論文タイトル・アブストラクト推敲サービス")
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))


def _static_version() -> str:
    files = [
        BASE_DIR / "static" / "app.js",
        BASE_DIR / "static" / "styles.css",
        BASE_DIR / "static" / "imagebuildabst.png",
    ]
    digest = hashlib.md5()
    for file_path in files:
        if file_path.exists():
            digest.update(file_path.read_bytes())
    return digest.hexdigest()[:8]


STATIC_VERSION = _static_version()


def _app_version() -> str:
    try:
        result = subprocess.run(
            ["git", "log", "-1", "--format=%cI"],
            cwd=BASE_DIR,
            capture_output=True,
            text=True,
            timeout=5,
            check=True,
        )
        commit_time = datetime.fromisoformat(result.stdout.strip())
        return commit_time.strftime("%Y-%m-%d %H:%M")
    except Exception:
        return "unknown"


APP_VERSION = _app_version()


class GenerationRequest(BaseModel):
    name: str = Field("", max_length=200)
    target_language: str = Field("ja", pattern="^(ja|en)$")
    background: str = Field("", max_length=1200)
    problem: str = Field("", max_length=1200)
    objective: str = Field("", max_length=1200)
    challenge: str = Field("", max_length=1200)
    methods: str = Field("", max_length=1200)
    experiment: str = Field("", max_length=1200)
    expected_results: str = Field("", max_length=1200)
    discussion: str = Field("", max_length=1200)
    conclusion: str = Field("", max_length=1200)
    contribution: str = Field("", max_length=1200)
    model: str = Field(DEFAULT_MODEL, min_length=1)


DEFAULT_ACADEMIC_FIELD = "情報工学全般（AI・データサイエンスを含む）"
DEFAULT_PAPER_TYPE = "journal article"
DEFAULT_TARGET_AUDIENCE = "学会査読者"
SECTION_DEFINITIONS: list[tuple[str, str]] = [
    ("background", "背景"),
    ("problem", "問題"),
    ("objective", "この論文での目的"),
    ("challenge", "目的達成のために解決すべき課題"),
    ("methods", "方法"),
    ("experiment", "実験"),
    ("expected_results", "想定される結果"),
    ("discussion", "課題が解決されているかの考察"),
    ("conclusion", "結論"),
    ("contribution", "貢献"),
]
SECTION_LABELS = dict(SECTION_DEFINITIONS)
SECTION_UPDATE_SCOPE_RULES: dict[str, str] = {
    "background": "背景だけを書く。問題提起、研究ギャップ、課題、目的、方法には踏み込まない。",
    "problem": "問題だけを書く。目的、課題、方法、結論には踏み込まない。",
    "objective": "この論文での目的だけを書く。課題や方法の具体的内容は入れない。",
    "challenge": "解決すべき課題だけを書く。具体的な方法や実験計画は入れない。",
    "methods": "方法だけを書く。実験条件や期待結果は入れない。",
    "experiment": "実験設計だけを書く。結果や考察は入れない。",
    "expected_results": "想定される結果だけを書く。考察や結論は入れない。",
    "discussion": "考察だけを書く。結論や貢献の断定までは入れない。",
    "conclusion": "結論だけを書く。貢献の詳述は入れない。",
    "contribution": "貢献だけを書く。背景や方法の説明に戻らない。",
}
EVALUATION_CRITERIA: list[dict[str, str]] = [
    {"key": "logic", "label": "全体の論理の説得性", "short": "論理"},
    {"key": "problem", "label": "問題の具体性", "short": "問題"},
    {"key": "objective", "label": "目的の具体性", "short": "目的"},
    {"key": "challenge", "label": "課題の具体性", "short": "課題"},
    {"key": "methods", "label": "方法の具体性", "short": "方法"},
    {"key": "experiment", "label": "実験の具体性", "short": "実験"},
    {"key": "expected_results", "label": "結果の定量性", "short": "結果"},
    {"key": "discussion", "label": "考察の具体性", "short": "考察"},
    {"key": "conclusion", "label": "結論の具体性", "short": "結論"},
    {"key": "contribution", "label": "貢献の具体性", "short": "貢献"},
]
EVALUATION_LABELS = {item["key"]: item["label"] for item in EVALUATION_CRITERIA}


def init_db() -> None:
    with closing(sqlite3.connect(DB_PATH)) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS generations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                research_topic TEXT NOT NULL,
                academic_field TEXT NOT NULL,
                paper_type TEXT NOT NULL,
                target_language TEXT NOT NULL,
                target_audience TEXT NOT NULL,
                objective TEXT NOT NULL,
                methods TEXT NOT NULL,
                key_points TEXT NOT NULL,
                constraints TEXT NOT NULL,
                source_title TEXT NOT NULL,
                source_text TEXT NOT NULL,
                guidance_text TEXT NOT NULL,
                model TEXT NOT NULL,
                recommended_title TEXT NOT NULL,
                abstract_text TEXT NOT NULL,
                result_json TEXT NOT NULL,
                cost_usd REAL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS generation_jobs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                status TEXT NOT NULL,
                progress INTEGER NOT NULL,
                stage TEXT NOT NULL,
                research_topic TEXT NOT NULL,
                source_filename TEXT NOT NULL,
                generation_id INTEGER,
                error_message TEXT NOT NULL,
                result_json TEXT NOT NULL
            )
            """
        )
        try:
            conn.execute("ALTER TABLE generations ADD COLUMN request_json TEXT NOT NULL DEFAULT ''")
        except sqlite3.OperationalError:
            pass
        try:
            conn.execute("ALTER TABLE generations ADD COLUMN name TEXT NOT NULL DEFAULT ''")
        except sqlite3.OperationalError:
            pass
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS cost_log (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                source TEXT NOT NULL,
                cost_usd REAL NOT NULL
            )
            """
        )
        conn.commit()


def fetch_and_log_cost(generation_id: str, api_key: str, source: str) -> None:
    cost_usd = fetch_openrouter_generation_cost(generation_id, api_key)
    log_cost(source, cost_usd)


def log_cost(source: str, cost_usd: float | None) -> None:
    if cost_usd is None:
        return
    with closing(sqlite3.connect(DB_PATH)) as conn:
        conn.execute(
            "INSERT INTO cost_log (created_at, source, cost_usd) VALUES (?, ?, ?)",
            (datetime.now(timezone.utc).isoformat(), source, cost_usd),
        )
        conn.commit()


init_db()


def fetch_models() -> list[str]:
    try:
        with httpx.Client(timeout=15) as client:
            response = client.get(f"{OLLAMA_URL}/api/tags")
            response.raise_for_status()
    except httpx.HTTPError as exc:
        raise HTTPException(status_code=503, detail=f"Ollama is unavailable: {exc}") from exc
    payload = response.json()
    return [item["name"] for item in payload.get("models", []) if item.get("name")]


def load_guidance() -> str:
    try:
        guidance_text = GUIDANCE_PATH.read_text(encoding="utf-8").strip()
    except FileNotFoundError as exc:
        raise HTTPException(status_code=500, detail=f"Guidance file not found: {GUIDANCE_PATH}") from exc

    if len(guidance_text) < 20:
        raise HTTPException(status_code=500, detail=f"Guidance file is too short: {GUIDANCE_PATH}")
    return guidance_text


def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(file_bytes))
    except Exception as exc:
        raise HTTPException(status_code=400, detail="PDF の読み取りに失敗しました。") from exc

    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(page.strip() for page in pages if page.strip()).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        document = Document(BytesIO(file_bytes))
    except Exception as exc:
        raise HTTPException(status_code=400, detail="DOCX の読み取りに失敗しました。") from exc

    texts: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if text:
            texts.append(text)

    seen_cells: set[int] = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_id = id(cell._tc)
                if cell_id in seen_cells:
                    continue
                seen_cells.add(cell_id)
                for para in cell.paragraphs:
                    text = para.text.strip()
                    if text:
                        texts.append(text)

    return "\n".join(texts).strip()


def extract_source_text_from_bytes(filename: str, file_bytes: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="補足資料は PDF または DOCX のみ対応しています。")
    if not file_bytes:
        raise HTTPException(status_code=400, detail="補足資料ファイルが空です。")

    if suffix == ".pdf":
        text = extract_text_from_pdf(file_bytes)
    else:
        text = extract_text_from_docx(file_bytes)

    if len(text) < 50:
        raise HTTPException(
            status_code=400,
            detail="補足資料から十分な本文を抽出できませんでした。文字ベースの PDF または DOCX を確認してください。",
        )
    return text[:MAX_SOURCE_TEXT_LENGTH]


def build_generation_prompt(request_data: GenerationRequest, guidance_text: str) -> str:
    language_label = "日本語" if request_data.target_language == "ja" else "English"
    filled_section_count = sum(1 for field_name, _label in SECTION_DEFINITIONS if getattr(request_data, field_name).strip())
    section_list_text = "\n".join(
        f"- {label}: {getattr(request_data, field_name) or '未入力'}"
        for field_name, label in SECTION_DEFINITIONS
    )
    evaluation_list_text = "\n".join(
        f'- criterion_key="{item["key"]}", criterion="{item["label"]}"'
        for item in EVALUATION_CRITERIA
    )
    section_update_rules = "\n".join(
        f'- field_name="{field_name}" ({label}): {SECTION_UPDATE_SCOPE_RULES[field_name]}'
        for field_name, label in SECTION_DEFINITIONS
    )
    return f"""
あなたは研究支援に強いアカデミック・ライティングアシスタントです。
目的は、与えられた短文メモをもとに、論文タイトルとアブストラクトを推敲することです。

出力は必ず妥当な JSON のみとし、Markdown のコードフェンスや余計な説明は含めないでください。

このサービスの前提:
- 対象分野は {DEFAULT_ACADEMIC_FIELD}。
- 論文種別は {DEFAULT_PAPER_TYPE}。
- 想定読者は {DEFAULT_TARGET_AUDIENCE}。

重要な制約:
- 出力言語は必ず {language_label} にする。
- 与えられていない実験結果、数値、評価指標、データセット名を捏造しない。
- 論文本文向けの断定的なトーンで書く。
- 情報が不足していても、過度に保留的な言い回しへ逃げず、入力から妥当な主張として自然にまとめる。
- 各項目について、問題点の指摘と、より論文らしい修正候補を返す。
- feedback は短く具体的に書く。
- revision_candidate は、元の意図をできるだけ保ちながら、学術文脈で自然な 1 文に整える。
- revised_title は簡潔で、焦点と貢献が伝わるものにする。
- revised_abstract は必ず 1 段落の連続した本文にする。
- revised_abstract に見出し、項目名、ラベル、箇条書き、番号付け、項目ごとの改行を入れない。
- revised_abstract では、「背景」「問題」「目的」「課題」「方法」「実験」「結果」「考察」「結論」「貢献」の順で内容を自然につなぐ。
- revised_abstract では、入力に含まれる情報を安易に要約、省略、統合しない。
- revised_abstract では、重要な条件、前提、制約、例外、数値、比較、評価、留保を落とさない。
- revised_abstract では、入力にない実験結果や数値を捏造せず、不足情報は断定的に作り足さない。
- keywords は 3 件から 6 件返す。
- overall_notes には、全体として直すとよい点を返す。
- next_steps には、次に埋めると完成度が上がる情報を返す。不要なら空配列でよい。
- section_feedback には、入力項目のうち内容があるものを順番通り返す。
- evaluation には、下記 10 項目をこの順番どおり必ず入れる。
- evaluation.score は 0 から 10 の整数にする。
- evaluation.comment には、その点数の理由と上げるための要点を短く具体的に書く。
- evaluation の採点対象は、生成された revised_title と revised_abstract そのものにする。
- ただし、その評価では「この最終生成リクエストで実際に使われた入力項目」を根拠にする。
- すでに項目推敲を採用して入力欄が更新されている場合は、その採用後の文を生成の素材とみなす。
- 項目推敲を採用していない場合は、その時点の元入力を生成の素材とみなす。
- つまり、採点は「現在の入力項目を素材として生成されたタイトル・アブストラクトが、どれだけ良くまとまっているか」を評価する。
- feedback、revision_candidate、section_updates は改善提案として返してよいが、それ自体を採点対象にはしない。
- 採点は「この時点の研究メモとしての完成度」を評価する。査読投稿直前の完成原稿基準で過度に厳しくしすぎない。
- 0 点から 1 点は、ほぼ空欄、意味不明、または論文要素として成立していない場合だけに使う。
- 2 点から 3 点は、断片的で意図は少し見えるが、論文要素としてはかなり不足している場合に使う。
- 4 点から 5 点は、短い研究メモとして意図は分かるが、具体性や説得力が不足している場合に使う。
- 6 点から 7 点は、短文メモとしては筋が通っており、推敲前として妥当だが、具体性や定量性がまだ足りない場合に使う。
- 8 点から 10 点は、項目としてかなり具体的で、論文アブストラクト要素として強い場合に使う。
- 入力に内容がある項目まで、機械的にすべて 0 点や 1 点へそろえない。
- すべての評価項目が本当に同程度に弱いと判断できる場合を除き、evaluation.score を一律同点にしない。
- section_updates には、10 項目すべてについて改善後の入力案を順番どおり返す。
- section_updates.suggested_text は、その項目の入力欄にそのまま戻せる短い学術文にする。
- section_updates.reason には、その修正案で何を補強したかを短く書く。
- section_updates では、各項目に別項目の内容を混ぜない。

JSON format:
{{
  "revised_title": "string",
  "revised_abstract": "string",
  "abstract_feedback": "string",
  "section_feedback": [
    {{
      "section": "string",
      "input_text": "string",
      "feedback": "string",
      "revision_candidate": "string"
    }}
  ],
  "evaluation": [
    {{
      "criterion_key": "string",
      "criterion": "string",
      "score": 0,
      "comment": "string"
    }}
  ],
  "section_updates": [
    {{
      "field_name": "string",
      "section": "string",
      "current_text": "string",
      "suggested_text": "string",
      "reason": "string"
    }}
  ],
  "keywords": ["string"],
  "overall_notes": ["string"],
  "next_steps": ["string"]
}}

evaluation の必須項目:
{evaluation_list_text}

section_updates の対象項目と制約:
{section_update_rules}

生成ガイド:
{guidance_text}

入力状況:
- 10 項目中 {filled_section_count} 項目に内容がある。
- 空欄の項目は低めに評価してよいが、内容がある項目は記述の具体性に応じて段階的に採点する。

入力項目:
{section_list_text}
""".strip()


def build_section_revision_prompt(request_data: GenerationRequest, field_name: str, guidance_text: str) -> str:
    language_label = "日本語" if request_data.target_language == "ja" else "English"
    section_label = SECTION_LABELS[field_name]
    current_text = getattr(request_data, field_name)
    other_sections = "\n".join(
        f"- {label}: {getattr(request_data, current_field) or '未入力'}"
        for current_field, label in SECTION_DEFINITIONS
        if current_field != field_name
    )
    excluded_sections = "、".join(label for current_field, label in SECTION_DEFINITIONS if current_field != field_name)
    section_specific_rules = ""
    if field_name == "background":
        section_specific_rules = """
- 背景の修正候補は、研究や社会の文脈説明だけに限定する。
- 背景の修正候補に、問題提起や研究ギャップを担う文を入れない。
- 「しかし」「一方で」「未解明」「未検証」「課題」「問題」など、問題・課題の提示につながる展開は `revision_candidate` に入れない。
- 背景で問題となる事象に触れたい場合は、`revision_candidate` ではなく `detected_problem` にだけ書く。
""".strip()
    elif field_name == "objective":
        section_specific_rules = """
- 目的の修正候補は、この論文で明らかにしたいことの説明だけに限定する。
- 目的の修正候補に、課題の具体的な中身や解決手段を入れない。
- 目的と背景・問題から、目的を達成するために解決すべき課題を推定し、`detected_challenge` に 1 文で返す。
- `detected_challenge` は目的本文に含めず、課題欄へ移すための補助情報としてのみ返す。
- 目的だけでは課題を確定しにくい場合でも、目的と他項目から最も妥当な課題を 1 つ補って返す。
""".strip()
    elif field_name == "problem":
        section_specific_rules = """
- 問題の修正候補は、背景の中で何が問題になっているかの説明だけに限定する。
- 問題の修正候補に、この論文で明らかにしたいこと（目的）そのものを書かない。
- 背景と問題から、この論文で明らかにしたいこと（目的）を推定し、`detected_objective` に 1 文で返す。
- `detected_objective` は問題本文に含めず、目的欄へ移すための補助情報としてのみ返す。
- 問題だけでは目的を確定しにくい場合でも、問題と他項目から最も妥当な目的を 1 つ補って返す。
""".strip()
    elif field_name == "challenge":
        section_specific_rules = """
- 課題の修正候補は、目的達成のために解決すべき課題の説明だけに限定する。
- 課題の修正候補に、具体的な解決手段（方法）そのものを書かない。
- 課題と目的・背景・問題から、課題を解決するために必要な方法を推定し、`detected_methods` に 1 文で返す。
- `detected_methods` は課題本文に含めず、方法欄へ移すための補助情報としてのみ返す。
- 課題だけでは方法を確定しにくい場合でも、課題と他項目から最も妥当な方法を 1 つ補って返す。
""".strip()
    elif field_name == "methods":
        section_specific_rules = """
- 方法の修正候補は、提案する手法やアプローチの説明だけに限定する。
- 方法の修正候補に、具体的な実験設計そのものを書かない。
- 方法と課題・目的から、その方法を検証するために必要な実験を推定し、`detected_experiment` に 1 文で返す。
- `detected_experiment` は方法本文に含めず、実験欄へ移すための補助情報としてのみ返す。
- 方法だけでは実験を確定しにくい場合でも、方法と他項目から最も妥当な実験を 1 つ補って返す。
""".strip()
    elif field_name == "experiment":
        section_specific_rules = """
- 実験の修正候補は、評価方法、比較条件、対象データの説明だけに限定する。
- 実験の修正候補に、想定される結果そのものを書かない。
- 実験と方法・課題から、その実験で想定される結果を推定し、`detected_expected_results` に 1 文で返す。
- `detected_expected_results` は実験本文に含めず、想定される結果欄へ移すための補助情報としてのみ返す。
- 実験だけでは想定される結果を確定しにくい場合でも、実験と他項目から最も妥当な想定される結果を 1 つ補って返す。
""".strip()
    elif field_name == "expected_results":
        section_specific_rules = """
- 想定される結果の修正候補は、期待される結果の説明だけに限定する。
- 想定される結果の修正候補に、課題が解決されているかどうかの考察そのものを書かない。
- 想定される結果と課題・目的から、その結果によって課題が解決されているかの考察を推定し、`detected_discussion` に 1 文で返す。
- `detected_discussion` は想定される結果本文に含めず、考察欄へ移すための補助情報としてのみ返す。
- 想定される結果だけでは考察を確定しにくい場合でも、想定される結果と他項目から最も妥当な考察を 1 つ補って返す。
""".strip()
    elif field_name == "discussion":
        section_specific_rules = """
- 考察の修正候補は、課題が解決されているかどうかの考察だけに限定する。
- 考察の修正候補に、結論そのものを書かない。
- 考察と目的・課題から、この論文の結論を推定し、`detected_conclusion` に 1 文で返す。
- `detected_conclusion` は考察本文に含めず、結論欄へ移すための補助情報としてのみ返す。
- 考察だけでは結論を確定しにくい場合でも、考察と他項目から最も妥当な結論を 1 つ補って返す。
""".strip()
    elif field_name == "conclusion":
        section_specific_rules = """
- 結論の修正候補は、この論文としての結び・まとめだけに限定する。
- 結論の修正候補に、貢献の具体的な中身を書かない。
- 結論と目的・課題から、この研究の貢献を推定し、`detected_contribution` に 1 文で返す。
- `detected_contribution` は結論本文に含めず、貢献欄へ移すための補助情報としてのみ返す。
- 結論だけでは貢献を確定しにくい場合でも、結論と他項目から最も妥当な貢献を 1 つ補って返す。
""".strip()
    return f"""
あなたは研究支援に強いアカデミック・ライティングアシスタントです。
目的は、論文アブストラクト構成要素の 1 項目だけを推敲することです。

出力は必ず妥当な JSON のみとし、Markdown のコードフェンスや余計な説明は含めないでください。

このサービスの前提:
- 対象分野は {DEFAULT_ACADEMIC_FIELD}。
- 論文種別は {DEFAULT_PAPER_TYPE}。
- 想定読者は {DEFAULT_TARGET_AUDIENCE}。
- 出力言語は必ず {language_label} にする。
- 論文本文向けの断定的なトーンで書く。

重要な制約:
- 与えられていない実験結果、数値、評価指標、データセット名を捏造しない。
- フィードバックは短く具体的に書く。
- 修正候補は 1 文から 2 文の自然な学術文にする。
- 修正候補は元の意図をできるだけ保つ。
- 今回は「{section_label}」だけを出力対象にする。
- {excluded_sections} の内容は書かない。
- 修正候補に、別項目の見出しや別項目に相当する内容を混ぜない。
- 例えば「背景」を推敲する場合、問題提起、研究目的、課題、方法、実験、結果、考察、結論、貢献までは踏み込まない。
- 「背景」を推敲する場合は、背景から自然に導ける問題となっている事象も推定し、`detected_problem` に 1 文で返す。
- `detected_problem` は背景本文に含めず、問題欄へ移すための補助情報としてのみ返す。
- 背景から複数の問題が考えられる場合は、この論文の主眼として最も自然な 1 つを選ぶ。
- 背景だけでは問題を確定しにくい場合でも、背景と他項目から最も妥当な問題を 1 つ補って返す。
{section_specific_rules}

JSON format:
{{
  "section": "string",
  "input_text": "string",
  "feedback": "string",
  "revision_candidate": "string",
  "detected_problem": "string",
  "detected_challenge": "string",
  "detected_objective": "string",
  "detected_methods": "string",
  "detected_experiment": "string",
  "detected_expected_results": "string",
  "detected_discussion": "string",
  "detected_conclusion": "string",
  "detected_contribution": "string"
}}

生成ガイド:
{guidance_text}

今回推敲する項目:
{section_label}

この項目の入力:
{current_text}

他の入力項目:
{other_sections}
""".strip()


def build_evaluation_prompt(request_data: GenerationRequest, revised_title: str, revised_abstract: str) -> str:
    language_label = "日本語" if request_data.target_language == "ja" else "English"
    section_list_text = "\n".join(
        f"- {label}: {getattr(request_data, field_name) or '未入力'}"
        for field_name, label in SECTION_DEFINITIONS
    )
    evaluation_list_text = "\n".join(
        f'- criterion_key="{item["key"]}", criterion="{item["label"]}"'
        for item in EVALUATION_CRITERIA
    )
    return f"""
あなたは研究支援に強いアカデミック・ライティングアシスタントです。
目的は、入力項目をもとに生成された論文タイトルとアブストラクトの完成度を評価することです。

出力は必ず妥当な JSON のみとし、Markdown のコードフェンスや余計な説明は含めないでください。

このサービスの前提:
- 対象分野は {DEFAULT_ACADEMIC_FIELD}。
- 論文種別は {DEFAULT_PAPER_TYPE}。
- 想定読者は {DEFAULT_TARGET_AUDIENCE}。
- 出力言語は必ず {language_label} にする。

重要な制約:
- evaluation の採点対象は、下記の generated title と generated abstract そのものにする。
- ただし、評価理由は、その生成結果が入力項目の意図をどれだけ適切に表現できているかに基づいて書く。
- 項目推敲をしたかどうかは採点基準に含めない。
- 0 点から 1 点は、該当要素が生成結果にほぼ現れていない、または意味不明な場合だけに使う。
- 2 点から 3 点は、要素は少し見えるが、かなり弱い場合に使う。
- 4 点から 5 点は、最低限は表現できているが、具体性や説得力が不足する場合に使う。
- 6 点から 7 点は、要素が概ね自然に表現されているが、精度や定量性がまだ弱い場合に使う。
- 8 点から 10 点は、要素が明確かつ説得的に表現されている場合に使う。
- 入力項目が未推敲でも、生成された title / abstract が十分にまとまっていれば低得点に固定しない。
- evaluation には、下記 10 項目をこの順番どおり必ず入れる。
- evaluation.score は 0 から 10 の整数にする。
- evaluation.comment には、その点数の理由と上げるための要点を短く具体的に書く。

JSON format:
{{
  "evaluation": [
    {{
      "criterion_key": "string",
      "criterion": "string",
      "score": 0,
      "comment": "string"
    }}
  ]
}}

evaluation の必須項目:
{evaluation_list_text}

generated title:
{revised_title or '生成なし'}

generated abstract:
{revised_abstract or '生成なし'}

入力項目:
{section_list_text}
""".strip()


def normalize_evaluation(raw_result: dict[str, Any]) -> list[dict[str, Any]]:
    normalized_evaluation: list[dict[str, Any]] = []
    evaluation_map: dict[str, dict[str, Any]] = {}
    for item in raw_result.get("evaluation", []):
        if not isinstance(item, dict):
            continue
        criterion_key = str(item.get("criterion_key", "")).strip()
        criterion = str(item.get("criterion", "")).strip()
        if not criterion_key and criterion:
            for meta in EVALUATION_CRITERIA:
                if meta["label"] == criterion:
                    criterion_key = meta["key"]
                    break
        if criterion_key not in EVALUATION_LABELS:
            continue
        try:
            score = int(item.get("score", 0))
        except (TypeError, ValueError):
            score = 0
        evaluation_map[criterion_key] = {
            "criterion_key": criterion_key,
            "criterion": EVALUATION_LABELS[criterion_key],
            "score": max(0, min(10, score)),
            "comment": str(item.get("comment", "")).strip(),
        }

    for meta in EVALUATION_CRITERIA:
        normalized_evaluation.append(
            evaluation_map.get(
                meta["key"],
                {
                    "criterion_key": meta["key"],
                    "criterion": meta["label"],
                    "score": 0,
                    "comment": "",
                },
            )
        )
    return normalized_evaluation


def apply_evaluation_result(result: dict[str, Any], evaluation: list[dict[str, Any]]) -> dict[str, Any]:
    pass_threshold = PASS_THRESHOLD
    result["evaluation"] = evaluation
    result["pass_threshold"] = pass_threshold
    result["pass_all"] = all(item["score"] >= pass_threshold for item in evaluation)
    return result


def normalize_generation(raw_result: dict[str, Any], request_data: GenerationRequest) -> dict[str, Any]:
    normalized_sections: list[dict[str, str]] = []
    for item in raw_result.get("section_feedback", []):
        if not isinstance(item, dict):
            continue
        section = str(item.get("section", "")).strip()
        input_text = str(item.get("input_text", "")).strip()
        feedback = str(item.get("feedback", "")).strip()
        revision_candidate = str(item.get("revision_candidate", "")).strip()
        if section:
            normalized_sections.append(
                {
                    "section": section,
                    "input_text": input_text,
                    "feedback": feedback,
                    "revision_candidate": revision_candidate,
                }
            )

    normalized_evaluation = normalize_evaluation(raw_result)

    feedback_by_section = {item["section"]: item for item in normalized_sections}
    updates_map: dict[str, dict[str, str]] = {}
    for item in raw_result.get("section_updates", []):
        if not isinstance(item, dict):
            continue
        field_name = str(item.get("field_name", "")).strip()
        if field_name not in SECTION_LABELS:
            section_label = str(item.get("section", "")).strip()
            for current_field, label in SECTION_DEFINITIONS:
                if section_label == label:
                    field_name = current_field
                    break
        if field_name not in SECTION_LABELS:
            continue
        suggested_text = str(item.get("suggested_text", "")).strip()
        if field_name == "background":
            suggested_text = sanitize_background_revision(suggested_text)
        updates_map[field_name] = {
            "field_name": field_name,
            "section": SECTION_LABELS[field_name],
            "current_text": str(item.get("current_text", "")).strip() or getattr(request_data, field_name),
            "suggested_text": suggested_text,
            "reason": str(item.get("reason", "")).strip(),
        }

    normalized_updates: list[dict[str, str]] = []
    for field_name, label in SECTION_DEFINITIONS:
        fallback_feedback = feedback_by_section.get(label, {})
        suggested_text = updates_map.get(field_name, {}).get("suggested_text", "").strip()
        if not suggested_text:
            suggested_text = str(fallback_feedback.get("revision_candidate", "")).strip() or getattr(request_data, field_name)
        if field_name == "background":
            suggested_text = sanitize_background_revision(suggested_text)
        normalized_updates.append(
            {
                "field_name": field_name,
                "section": label,
                "current_text": updates_map.get(field_name, {}).get("current_text", "").strip() or getattr(request_data, field_name),
                "suggested_text": suggested_text,
                "reason": updates_map.get(field_name, {}).get("reason", "").strip()
                or str(fallback_feedback.get("feedback", "")).strip(),
            }
        )

    revised_title = str(raw_result.get("revised_title", "")).strip()
    revised_abstract = str(raw_result.get("revised_abstract", "")).strip()
    result = {
        "revised_title": revised_title or "タイトル案を生成できませんでした",
        "revised_abstract": revised_abstract,
        "abstract_feedback": str(raw_result.get("abstract_feedback", "")).strip(),
        "section_feedback": normalized_sections,
        "evaluation": normalized_evaluation,
        "pass_threshold": PASS_THRESHOLD,
        "pass_all": all(item["score"] >= PASS_THRESHOLD for item in normalized_evaluation),
        "section_updates": normalized_updates,
        "keywords": [str(item).strip() for item in raw_result.get("keywords", []) if str(item).strip()][:6],
        "overall_notes": [str(item).strip() for item in raw_result.get("overall_notes", []) if str(item).strip()][:6],
        "next_steps": [str(item).strip() for item in raw_result.get("next_steps", []) if str(item).strip()][:6],
    }
    return result


def sanitize_background_revision(text: str) -> str:
    cleaned = re.sub(r"\s+", " ", text).strip()
    if not cleaned:
        return cleaned

    split_markers = ("しかし", "しかしながら", "一方で", "ただし", "ただ", "なお")
    for marker in split_markers:
        marker_index = cleaned.find(marker)
        if marker_index > 0:
            candidate = cleaned[:marker_index].rstrip("、 ,")
            if candidate:
                return candidate + ("。" if not candidate.endswith("。") else "")

    problem_cues = ("未解明", "未検証", "課題", "問題", "不足", "困難")
    sentences = [item.strip() for item in re.split(r"(?<=[。！？])", cleaned) if item.strip()]
    if len(sentences) >= 2:
        filtered: list[str] = []
        for index, sentence in enumerate(sentences):
            if index > 0 and any(cue in sentence for cue in problem_cues):
                break
            filtered.append(sentence)
        if filtered:
            return "".join(filtered).strip()

    return cleaned


DETECTED_NEXT_FIELD_KEY: dict[str, str] = {
    "background": "detected_problem",
    "problem": "detected_objective",
    "objective": "detected_challenge",
    "challenge": "detected_methods",
    "methods": "detected_experiment",
    "experiment": "detected_expected_results",
    "expected_results": "detected_discussion",
    "discussion": "detected_conclusion",
    "conclusion": "detected_contribution",
}


def normalize_section_revision(raw_result: dict[str, Any], field_name: str, input_text: str) -> dict[str, str]:
    revision_candidate = str(raw_result.get("revision_candidate", "")).strip()
    if field_name == "background":
        revision_candidate = sanitize_background_revision(revision_candidate)
    result = {
        "field_name": field_name,
        "section": str(raw_result.get("section", "")).strip() or SECTION_LABELS[field_name],
        "input_text": str(raw_result.get("input_text", "")).strip() or input_text,
        "feedback": str(raw_result.get("feedback", "")).strip(),
        "revision_candidate": revision_candidate,
    }
    detected_key = DETECTED_NEXT_FIELD_KEY.get(field_name)
    if detected_key:
        result[detected_key] = str(raw_result.get(detected_key, "")).strip()
    return result


def parse_json_response(raw_response: str) -> dict[str, Any]:
    cleaned = raw_response.strip()
    if not cleaned:
        raise HTTPException(status_code=502, detail="AI が空の応答を返しました。")

    if cleaned.startswith("```"):
        cleaned = cleaned.strip("`")
        if cleaned.startswith("json"):
            cleaned = cleaned[4:].strip()

    try:
        parsed = json.loads(cleaned)
    except json.JSONDecodeError:
        decoder = json.JSONDecoder()
        start = cleaned.find("{")
        if start == -1:
            raise HTTPException(status_code=502, detail="AI の JSON 応答を解析できませんでした。")
        try:
            parsed, _ = decoder.raw_decode(cleaned[start:])
        except json.JSONDecodeError as exc:
            raise HTTPException(status_code=502, detail="AI の JSON 応答を解析できませんでした。") from exc

    if not isinstance(parsed, dict):
        raise HTTPException(status_code=502, detail="AI の応答形式が不正です。")
    return parsed


def fetch_openrouter_generation_cost(
    generation_id: str, api_key: str, max_attempts: int = 12, retry_delay: float = 2.0
) -> float | None:
    # OpenRouter's cost stats for a generation are not indexed immediately after
    # the completion response returns, so an immediate lookup routinely 404s.
    for attempt in range(max_attempts):
        try:
            with httpx.Client(timeout=10) as client:
                response = client.get(
                    f"{_OPENROUTER_BASE}/generation",
                    params={"id": generation_id},
                    headers={"Authorization": f"Bearer {api_key}"},
                )
                response.raise_for_status()
                return float(response.json().get("data", {}).get("total_cost") or 0)
        except Exception:
            if attempt < max_attempts - 1:
                time.sleep(retry_delay)
                continue
            return None
    return None


def fetch_openrouter_credits(api_key: str) -> dict[str, Any] | None:
    try:
        with httpx.Client(timeout=10) as client:
            response = client.get(
                f"{_OPENROUTER_BASE}/credits",
                headers={"Authorization": f"Bearer {api_key}"},
            )
            response.raise_for_status()
            data = response.json().get("data", {})
            total = float(data.get("total_credits") or 0)
            usage = float(data.get("total_usage") or 0)
            return {"total_credits": total, "usage": usage, "remaining": round(total - usage, 6)}
    except Exception:
        return None


def generate_openrouter(
    request_data: GenerationRequest,
    guidance_text: str,
    job_id: int | None = None,
    api_key: str = "",
) -> tuple[dict[str, Any], float | None]:
    prompt = build_generation_prompt(request_data, guidance_text)
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "HTTP-Referer": "http://localhost:8000",
        "X-Title": "BuildAbstract",
    }
    payload = {
        "model": OPENROUTER_MODEL,
        "messages": [{"role": "user", "content": prompt}],
        "stream": True,
        "temperature": 0.4,
    }

    if job_id is not None:
        _llm_stream[job_id] = ""

    raw_response = ""
    generation_id: str | None = None
    deadline = time.monotonic() + REQUEST_TIMEOUT
    try:
        with httpx.Client(timeout=httpx.Timeout(connect=30, read=REQUEST_TIMEOUT, write=30, pool=10)) as client:
            with client.stream("POST", f"{_OPENROUTER_BASE}/chat/completions", json=payload, headers=headers) as response:
                response.raise_for_status()
                for line in response.iter_lines():
                    if time.monotonic() > deadline:
                        raise HTTPException(status_code=503, detail="生成がタイムアウトしました（処理時間超過）。")
                    if not line or not line.startswith("data: "):
                        continue
                    data = line[6:]
                    if data.strip() == "[DONE]":
                        break
                    try:
                        chunk = json.loads(data)
                    except json.JSONDecodeError:
                        continue
                    if generation_id is None:
                        generation_id = chunk.get("id")
                    token = (chunk.get("choices") or [{}])[0].get("delta", {}).get("content") or ""
                    raw_response += token
                    if job_id is not None:
                        _llm_stream[job_id] = raw_response
    except httpx.HTTPError as exc:
        raise HTTPException(status_code=503, detail=f"OpenRouter での推敲に失敗しました: {exc}") from exc

    if job_id is not None:
        update_job(job_id, progress=90, stage="コストを集計しています。")
    cost_usd = fetch_openrouter_generation_cost(generation_id, api_key) if generation_id else None
    parsed = parse_json_response(raw_response.strip())
    result = normalize_generation(parsed, request_data)
    if job_id is not None:
        update_job(job_id, progress=94, stage="生成されたタイトルとアブストラクトを評価しています。")
    evaluation = evaluate_generated_outputs_openrouter(request_data, result, api_key)
    return apply_evaluation_result(result, evaluation), cost_usd


def generate_ollama(
    request_data: GenerationRequest,
    guidance_text: str,
    job_id: int | None = None,
    ollama_url: str | None = None,
) -> dict[str, Any]:
    prompt = build_generation_prompt(request_data, guidance_text)
    target_url = ollama_url or OLLAMA_URLS[0]
    payload = {
        "model": request_data.model,
        "prompt": prompt,
        "stream": True,
        "options": {"temperature": 0.4},
    }

    if job_id is not None:
        _llm_stream[job_id] = ""

    raw_response = ""
    deadline = time.monotonic() + REQUEST_TIMEOUT
    try:
        with httpx.Client(timeout=httpx.Timeout(connect=30, read=REQUEST_TIMEOUT, write=30, pool=10)) as client:
            with client.stream("POST", f"{target_url}/api/generate", json=payload) as response:
                response.raise_for_status()
                for line in response.iter_lines():
                    if time.monotonic() > deadline:
                        raise HTTPException(status_code=503, detail="生成がタイムアウトしました（処理時間超過）。")
                    if not line:
                        continue
                    try:
                        chunk = json.loads(line)
                    except json.JSONDecodeError:
                        continue
                    token = chunk.get("response", "")
                    raw_response += token
                    if job_id is not None:
                        _llm_stream[job_id] = raw_response
                    if chunk.get("done"):
                        break
    except httpx.HTTPError as exc:
        raise HTTPException(status_code=503, detail=f"Ollama での推敲に失敗しました: {exc}") from exc

    parsed = parse_json_response(raw_response.strip())
    result = normalize_generation(parsed, request_data)
    if job_id is not None:
        update_job(job_id, progress=94, stage="生成されたタイトルとアブストラクトを評価しています。")
    evaluation = evaluate_generated_outputs_ollama(request_data, result, target_url)
    return apply_evaluation_result(result, evaluation)


def evaluate_generated_outputs_openrouter(
    request_data: GenerationRequest,
    result: dict[str, Any],
    api_key: str,
) -> list[dict[str, Any]]:
    prompt = build_evaluation_prompt(
        request_data,
        str(result.get("revised_title", "")).strip(),
        str(result.get("revised_abstract", "")).strip(),
    )
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "HTTP-Referer": "http://localhost:8000",
        "X-Title": "BuildAbstract",
    }
    payload = {
        "model": OPENROUTER_MODEL,
        "messages": [{"role": "user", "content": prompt}],
        "stream": False,
        "temperature": 0.2,
    }
    try:
        with httpx.Client(timeout=httpx.Timeout(connect=30, read=REQUEST_TIMEOUT, write=30, pool=10)) as client:
            response = client.post(f"{_OPENROUTER_BASE}/chat/completions", json=payload, headers=headers)
            response.raise_for_status()
    except httpx.HTTPError as exc:
        raise HTTPException(status_code=503, detail=f"OpenRouter での評価に失敗しました: {exc}") from exc

    response_json = response.json()
    content = (((response_json.get("choices") or [{}])[0].get("message") or {}).get("content") or "")
    parsed = parse_json_response(content)
    return normalize_evaluation(parsed)


def evaluate_generated_outputs_ollama(
    request_data: GenerationRequest,
    result: dict[str, Any],
    ollama_url: str,
) -> list[dict[str, Any]]:
    prompt = build_evaluation_prompt(
        request_data,
        str(result.get("revised_title", "")).strip(),
        str(result.get("revised_abstract", "")).strip(),
    )
    payload = {
        "model": request_data.model,
        "prompt": prompt,
        "stream": False,
        "options": {"temperature": 0.2},
    }
    try:
        with httpx.Client(timeout=httpx.Timeout(connect=30, read=REQUEST_TIMEOUT, write=30, pool=10)) as client:
            response = client.post(f"{ollama_url}/api/generate", json=payload)
            response.raise_for_status()
    except httpx.HTTPError as exc:
        raise HTTPException(status_code=503, detail=f"Ollama での評価に失敗しました: {exc}") from exc

    parsed = parse_json_response((response.json().get("response") or "").strip())
    return normalize_evaluation(parsed)


def revise_section_openrouter(
    request_data: GenerationRequest,
    field_name: str,
    guidance_text: str,
    api_key: str,
) -> tuple[dict[str, str], str | None]:
    prompt = build_section_revision_prompt(request_data, field_name, guidance_text)
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "HTTP-Referer": "http://localhost:8000",
        "X-Title": "BuildAbstract",
    }
    payload = {
        "model": OPENROUTER_MODEL,
        "messages": [{"role": "user", "content": prompt}],
        "stream": False,
        "temperature": 0.35,
    }
    try:
        with httpx.Client(timeout=httpx.Timeout(connect=30, read=REQUEST_TIMEOUT, write=30, pool=10)) as client:
            response = client.post(f"{_OPENROUTER_BASE}/chat/completions", json=payload, headers=headers)
            response.raise_for_status()
    except httpx.HTTPError as exc:
        raise HTTPException(status_code=503, detail=f"OpenRouter での項目推敲に失敗しました: {exc}") from exc

    response_json = response.json()
    content = (
        ((response_json.get("choices") or [{}])[0].get("message") or {}).get("content") or ""
    )
    generation_id = response_json.get("id")
    parsed = parse_json_response(content)
    return normalize_section_revision(parsed, field_name, getattr(request_data, field_name)), generation_id


def revise_section_ollama(
    request_data: GenerationRequest,
    field_name: str,
    guidance_text: str,
    ollama_url: str,
) -> dict[str, str]:
    prompt = build_section_revision_prompt(request_data, field_name, guidance_text)
    payload = {
        "model": request_data.model,
        "prompt": prompt,
        "stream": False,
        "options": {"temperature": 0.35},
    }
    try:
        with httpx.Client(timeout=httpx.Timeout(connect=30, read=REQUEST_TIMEOUT, write=30, pool=10)) as client:
            response = client.post(f"{ollama_url}/api/generate", json=payload)
            response.raise_for_status()
    except httpx.HTTPError as exc:
        raise HTTPException(status_code=503, detail=f"Ollama での項目推敲に失敗しました: {exc}") from exc

    parsed = parse_json_response((response.json().get("response") or "").strip())
    return normalize_section_revision(parsed, field_name, getattr(request_data, field_name))


def save_generation(
    request_data: GenerationRequest,
    guidance_text: str,
    result: dict[str, Any],
    cost_usd: float | None = None,
) -> int:
    with closing(sqlite3.connect(DB_PATH)) as conn:
        cursor = conn.execute(
            """
            INSERT INTO generations (
                created_at, research_topic, academic_field, paper_type, target_language,
                target_audience, objective, methods, key_points, constraints,
                source_title, source_text, guidance_text, model,
                recommended_title, abstract_text, result_json, cost_usd, request_json, name
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                datetime.now(timezone.utc).isoformat(),
                result["revised_title"],
                DEFAULT_ACADEMIC_FIELD,
                DEFAULT_PAPER_TYPE,
                request_data.target_language,
                DEFAULT_TARGET_AUDIENCE,
                request_data.objective,
                request_data.methods,
                request_data.contribution,
                "",
                "",
                "",
                guidance_text,
                request_data.model,
                result["revised_title"],
                result["revised_abstract"],
                json.dumps(result, ensure_ascii=False),
                cost_usd,
                json.dumps(request_data.model_dump(), ensure_ascii=False),
                request_data.name,
            ),
        )
        conn.commit()
        return int(cursor.lastrowid)


def save_rescored_generation_result(generation_id: int, rescored_result: dict[str, Any]) -> None:
    with closing(sqlite3.connect(DB_PATH)) as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute(
            "SELECT result_json FROM generations WHERE id = ?",
            (generation_id,),
        ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Generation not found.")

        stored_result = json.loads(row["result_json"]) if row["result_json"] else {}
        stored_result.update(
            {
                "revised_title": rescored_result.get("revised_title", stored_result.get("revised_title", "")),
                "revised_abstract": rescored_result.get("revised_abstract", stored_result.get("revised_abstract", "")),
                "evaluation": rescored_result.get("evaluation", stored_result.get("evaluation", [])),
                "pass_threshold": rescored_result.get("pass_threshold", stored_result.get("pass_threshold", PASS_THRESHOLD)),
                "pass_all": rescored_result.get("pass_all", stored_result.get("pass_all", False)),
            }
        )

        conn.execute(
            """
            UPDATE generations
            SET research_topic = ?, recommended_title = ?, abstract_text = ?, result_json = ?
            WHERE id = ?
            """,
            (
                str(stored_result.get("revised_title", "")).strip(),
                str(stored_result.get("revised_title", "")).strip(),
                str(stored_result.get("revised_abstract", "")).strip(),
                json.dumps(stored_result, ensure_ascii=False),
                generation_id,
            ),
        )
        conn.commit()


def list_recent_generations(limit: int = 10) -> list[dict[str, Any]]:
    with closing(sqlite3.connect(DB_PATH)) as conn:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            """
            SELECT id, created_at, research_topic, academic_field, target_language,
                   model, recommended_title
            FROM generations
            ORDER BY id DESC
            LIMIT ?
            """,
            (limit,),
        ).fetchall()

    return [dict(row) for row in rows]


def create_job(research_topic: str, source_filename: str) -> int:
    now = datetime.now(timezone.utc).isoformat()
    with closing(sqlite3.connect(DB_PATH)) as conn:
        cursor = conn.execute(
            """
            INSERT INTO generation_jobs (
                created_at, updated_at, status, progress, stage,
                research_topic, source_filename, generation_id, error_message, result_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                now,
                now,
                "queued",
                5,
                "入力内容を受け付けました。",
                research_topic,
                source_filename,
                None,
                "",
                "",
            ),
        )
        conn.commit()
        return int(cursor.lastrowid)


def update_job(
    job_id: int,
    *,
    status: str | None = None,
    progress: int | None = None,
    stage: str | None = None,
    generation_id: int | None = None,
    error_message: str | None = None,
    result_payload: dict[str, Any] | None = None,
) -> None:
    updates = ["updated_at = ?"]
    params: list[Any] = [datetime.now(timezone.utc).isoformat()]

    if status is not None:
        updates.append("status = ?")
        params.append(status)
    if progress is not None:
        updates.append("progress = ?")
        params.append(progress)
    if stage is not None:
        updates.append("stage = ?")
        params.append(stage)
    if generation_id is not None:
        updates.append("generation_id = ?")
        params.append(generation_id)
    if error_message is not None:
        updates.append("error_message = ?")
        params.append(error_message)
    if result_payload is not None:
        updates.append("result_json = ?")
        params.append(json.dumps(result_payload, ensure_ascii=False))

    params.append(job_id)
    with closing(sqlite3.connect(DB_PATH)) as conn:
        conn.execute(f"UPDATE generation_jobs SET {', '.join(updates)} WHERE id = ?", params)
        conn.commit()


def get_job(job_id: int) -> dict[str, Any]:
    with closing(sqlite3.connect(DB_PATH)) as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute(
            """
            SELECT id, created_at, updated_at, status, progress, stage, research_topic,
                   source_filename, generation_id, error_message, result_json
            FROM generation_jobs
            WHERE id = ?
            """,
            (job_id,),
        ).fetchone()

    if row is None:
        raise HTTPException(status_code=404, detail="Job not found.")

    payload = dict(row)
    result = json.loads(payload["result_json"]) if payload["result_json"] else None
    payload["result"] = result
    payload.pop("result_json", None)
    live_output = _llm_stream.get(job_id, "")
    stored_output = (result or {}).get("llm_output", "") if result else ""
    payload["llm_output"] = live_output or stored_output
    payload["timeout_seconds"] = int(REQUEST_TIMEOUT)
    if USE_OPENROUTER:
        payload["pool_size"] = OPENROUTER_MAX_CONCURRENCY
        payload["pool_available"] = _openrouter_pool.qsize()
    else:
        payload["pool_size"] = len(OLLAMA_URLS)
        payload["pool_available"] = _ollama_pool.qsize()
    return payload


def get_generation_record(generation_id: int) -> dict[str, Any]:
    with closing(sqlite3.connect(DB_PATH)) as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute(
            """
            SELECT id, created_at, research_topic, academic_field, paper_type, target_language,
                   target_audience, objective, methods, key_points, constraints,
                   source_title, source_text, guidance_text, model,
                   recommended_title, abstract_text, result_json, cost_usd, request_json, name
            FROM generations
            WHERE id = ?
            """,
            (generation_id,),
        ).fetchone()

    if row is None:
        raise HTTPException(status_code=404, detail="Generation not found.")
    return dict(row)


def build_generation_markdown(record: dict[str, Any]) -> str:
    result = json.loads(record["result_json"])
    request_payload = json.loads(record.get("request_json") or "{}")
    evaluation = result.get("evaluation", [])
    pass_threshold = float(result.get("pass_threshold", PASS_THRESHOLD) or PASS_THRESHOLD)
    section_updates = result.get("section_updates", [])
    lines = [
        f"# {record['research_topic']} 推敲結果",
        "",
        "## 基本情報",
        f"- 生成ID: {record['id']}",
        f"- 作成者: {record.get('name') or '未入力'}",
        f"- 作成日時: {record['created_at']}",
        f"- 分野: {DEFAULT_ACADEMIC_FIELD}",
        f"- 論文種別: {DEFAULT_PAPER_TYPE}",
        f"- 出力言語: {'日本語' if record['target_language'] == 'ja' else 'English'}",
        f"- 想定読者: {DEFAULT_TARGET_AUDIENCE}",
        f"- 利用モデル: {record['model']}",
        "",
        "## 改稿タイトル",
        result.get("revised_title", "") or "生成なし",
        "",
        "## 項目別フィードバック",
    ]

    section_feedback = result.get("section_feedback", [])
    if section_feedback:
        for item in section_feedback:
            lines.extend(
                [
                    f"### {item.get('section', '項目')}",
                    f"- 入力: {item.get('input_text', '') or '未入力'}",
                    f"- フィードバック: {item.get('feedback', '') or 'コメントなし'}",
                    f"- 修正候補: {item.get('revision_candidate', '') or '候補なし'}",
                    "",
                ]
            )
    else:
        lines.extend(["- フィードバックなし", ""])

    lines.extend(
        [
            "## 改稿アブストラクト",
            result.get("revised_abstract", "") or "生成なし",
            "",
            "## アブストラクトへのコメント",
            result.get("abstract_feedback", "") or "コメントなし",
            "",
        ]
    )

    lines.extend(
        [
            "## 評価スコア",
            f"- 合格条件: 全項目 {pass_threshold:g} 点以上",
            f"- 判定: {'合格' if result.get('pass_all') else '未到達'}",
        ]
    )
    if evaluation:
        for item in evaluation:
            lines.append(
                f"- {item.get('criterion', '評価項目')}: {item.get('score', 0)}/10"
                f" ({item.get('comment', '') or 'コメントなし'})"
            )
    else:
        lines.append("- 評価なし")
    lines.append("")

    lines.append("## 改善用の更新案")
    if section_updates:
        for item in section_updates:
            lines.extend(
                [
                    f"### {item.get('section', '項目')}",
                    f"- 現在の入力: {item.get('current_text', '') or '未入力'}",
                    f"- 更新案: {item.get('suggested_text', '') or '候補なし'}",
                    f"- 補強ポイント: {item.get('reason', '') or 'コメントなし'}",
                    "",
                ]
            )
    else:
        lines.extend(["- 更新案なし", ""])

    def add_list_section(title: str, items: list[str]) -> None:
        lines.append(f"## {title}")
        if items:
            for item in items:
                lines.append(f"- {item}")
        else:
            lines.append("- 該当なし")
        lines.append("")

    add_list_section("キーワード", result.get("keywords", []))
    add_list_section("全体メモ", result.get("overall_notes", []))
    add_list_section("次に詰めるとよい点", result.get("next_steps", []))

    lines.extend(
        [
            "## 入力内容",
            f"- 背景: {request_payload.get('background') or '未入力'}",
            f"- 問題: {request_payload.get('problem') or '未入力'}",
            f"- この論文での目的: {request_payload.get('objective') or record['objective'] or '未入力'}",
            f"- 目的達成のために解決すべき課題: {request_payload.get('challenge') or '未入力'}",
            f"- 方法: {request_payload.get('methods') or record['methods'] or '未入力'}",
            f"- 実験: {request_payload.get('experiment') or '未入力'}",
            f"- 想定される結果: {request_payload.get('expected_results') or '未入力'}",
            f"- 課題が解決されているかの考察: {request_payload.get('discussion') or '未入力'}",
            f"- 結論: {request_payload.get('conclusion') or '未入力'}",
            f"- 貢献: {request_payload.get('contribution') or record['key_points'] or '未入力'}",
            "",
            "## 推敲ガイド",
            record["guidance_text"],
        ]
    )
    return "\n".join(lines).strip() + "\n"


def build_synology_chat_message(record: dict[str, Any]) -> str:
    result = json.loads(record["result_json"])
    evaluation = result.get("evaluation", [])
    total_score = sum(int(item.get("score", 0) or 0) for item in evaluation)
    lines = [
        "@m-ochi 確認をお願いします",
        "",
        f"名前: {record.get('name') or '未入力'}",
        f"タイトル: {result.get('revised_title', '') or '生成なし'}",
        "アブストラクト:",
        result.get("revised_abstract", "") or "生成なし",
        "",
    ]
    if evaluation:
        lines.append(f"総合点: {total_score}/{len(evaluation) * 10}")
        for item in evaluation:
            lines.append(f"{item.get('criterion', '評価項目')}: {int(item.get('score', 0) or 0)}/10")
    return "\n".join(lines).strip()


SYNOLOGY_CHAT_MESSAGE_LIMIT = 2000


def _hard_split(text: str, limit: int) -> list[str]:
    return [text[i : i + limit] for i in range(0, len(text), limit)]


def split_chat_message(text: str, limit: int = SYNOLOGY_CHAT_MESSAGE_LIMIT) -> list[str]:
    if len(text) <= limit:
        return [text]

    chunks: list[str] = []
    current = ""

    def flush() -> None:
        nonlocal current
        if current:
            chunks.append(current)
            current = ""

    for line in text.split("\n"):
        candidate = f"{current}\n{line}" if current else line
        if len(candidate) <= limit:
            current = candidate
            continue

        flush()
        if len(line) <= limit:
            current = line
            continue

        # A single line still exceeds the limit; fall back to splitting it at
        # sentence-ending punctuation (or hard-splitting) before continuing.
        for sentence in (s for s in re.split(r"(?<=[。！？])", line) if s):
            candidate = f"{current}{sentence}" if current else sentence
            if len(candidate) <= limit:
                current = candidate
                continue
            flush()
            if len(sentence) <= limit:
                current = sentence
            else:
                pieces = _hard_split(sentence, limit)
                chunks.extend(pieces[:-1])
                current = pieces[-1]

    flush()
    return chunks


SYNOLOGY_CHAT_POST_INTERVAL = 1.5


def _post_synology_chat_chunk(chunk: str, max_attempts: int = 3) -> None:
    error_message = "Synology Chat が投稿を受け付けませんでした。"
    for attempt in range(max_attempts):
        try:
            response = httpx.post(
                SYNOLOGY_CHAT_WEBHOOK_URL,
                data={"payload": json.dumps({"text": chunk}, ensure_ascii=False)},
                timeout=20.0,
            )
            response.raise_for_status()
            payload = response.json()
            if payload.get("success"):
                return
            error_message = payload.get("error", {}).get("errors") or error_message
        except httpx.HTTPError as exc:
            error_message = str(exc)

        if attempt < max_attempts - 1:
            # Synology Chat rate-limits posts made in quick succession
            # ("create post too fast"); back off and retry.
            time.sleep(SYNOLOGY_CHAT_POST_INTERVAL * (attempt + 2))
            continue
        raise HTTPException(status_code=502, detail=f"Synology Chat への投稿に失敗しました: {error_message}")


def post_synology_chat_message(message_text: str) -> None:
    if not synology_chat_configured():
        raise HTTPException(status_code=503, detail="Synology Chat の Webhook URL が未設定です。")

    for index, chunk in enumerate(split_chat_message(message_text)):
        if index > 0:
            time.sleep(SYNOLOGY_CHAT_POST_INTERVAL)
        _post_synology_chat_chunk(chunk)


def process_generation_job(
    job_id: int,
    *,
    backend_name: str,
    name: str,
    target_language: str,
    background: str,
    problem: str,
    objective: str,
    challenge: str,
    methods: str,
    experiment: str,
    expected_results: str,
    discussion: str,
    conclusion: str,
    contribution: str,
    model: str,
) -> None:
    try:
        update_job(job_id, status="processing", progress=10, stage="入力内容を確認しています。")
        update_job(job_id, progress=25, stage="フォーム入力を整理しています。")

        update_job(job_id, progress=40, stage="推敲ガイドを読み込んでいます。")
        guidance_text = load_guidance()

        request_data = GenerationRequest(
            name=name,
            target_language=target_language,
            background=background,
            problem=problem,
            objective=objective,
            challenge=challenge,
            methods=methods,
            experiment=experiment,
            expected_results=expected_results,
            discussion=discussion,
            conclusion=conclusion,
            contribution=contribution,
            model=model,
        )

        fixed_request = request_data.model_copy(update={"model": model_for_backend(backend_name)})
        update_job(job_id, progress=55, stage="推敲プロンプトを準備しています。")
        update_job(job_id, progress=65, stage="AI の処理順番を待っています。")

        if backend_name == "openrouter":
            active_pool = _openrouter_pool
            pool_slots = OPENROUTER_MAX_CONCURRENCY
            first_stage = f"OpenRouter ({OPENROUTER_MODEL}) で推敲しています。"
        else:
            active_pool = _ollama_pool
            pool_slots = len(OLLAMA_URLS)
            first_stage = "ローカル LLM で推敲しています。"

        max_attempts = min(3, max(1, pool_slots))
        last_exc: HTTPException | None = None
        result = None
        cost_usd: float | None = None

        for attempt in range(max_attempts):
            slot = active_pool.get()
            try:
                stage_message = first_stage if attempt == 0 else f"推敲を再試行しています（{attempt + 1}回目）。"
                update_job(job_id, progress=72, stage=stage_message)
                if backend_name == "openrouter":
                    result, cost_usd = generate_openrouter(fixed_request, guidance_text, job_id=job_id, api_key=slot)
                else:
                    result = generate_ollama(fixed_request, guidance_text, job_id=job_id, ollama_url=slot)
                break
            except HTTPException as exc:
                last_exc = exc
                if attempt < max_attempts - 1 and exc.status_code == 503:
                    logging.warning("generate attempt %d failed: %s", attempt + 1, exc.detail)
                    time.sleep(2)
                else:
                    raise
            finally:
                active_pool.put(slot)

        if result is None:
            raise last_exc or HTTPException(status_code=503, detail="推敲に失敗しました。")

        update_job(job_id, progress=88, stage="推敲結果を整形しています。")
        generation_id = save_generation(fixed_request, guidance_text, result, cost_usd=cost_usd)
        log_cost("generation", cost_usd)
        final_llm_output = _llm_stream.pop(job_id, "")
        result_payload = {
            "id": generation_id,
            "request": fixed_request.model_dump(),
            "result": result,
            "llm_output": final_llm_output,
            "cost_usd": cost_usd,
        }
        update_job(
            job_id,
            status="completed",
            progress=100,
            stage="タイトルとアブストラクトの推敲が完了しました。",
            generation_id=generation_id,
            result_payload=result_payload,
            error_message="",
        )
    except HTTPException as exc:
        message = exc.detail if isinstance(exc.detail, str) else "推敲中にエラーが発生しました。"
        update_job(job_id, status="failed", progress=100, stage="推敲に失敗しました。", error_message=message)
        _llm_stream.pop(job_id, None)
    except Exception as exc:
        logging.error("Unexpected error in generation job %d: %s", job_id, traceback.format_exc())
        update_job(
            job_id,
            status="failed",
            progress=100,
            stage="推敲に失敗しました。",
            error_message=f"予期しないエラーが発生しました: {exc}",
        )
        _llm_stream.pop(job_id, None)


@app.get("/api/debug/stream")
def debug_stream() -> dict[str, Any]:
    return {str(job_id): text[:120] for job_id, text in _llm_stream.items()}


@app.on_event("startup")
def startup() -> None:
    init_db()


@app.get("/", response_class=HTMLResponse)
def index(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request=request,
        name="index.html",
        context={
            "default_model": model_for_backend(current_backend()),
            "default_backend": current_backend(),
            "guidance_text": load_guidance(),
            "backend": current_backend(),
            "static_version": STATIC_VERSION,
            "app_version": APP_VERSION,
        },
    )


@app.get("/api/health")
def health() -> dict[str, Any]:
    try:
        models = fetch_models()
        ollama_model_available = DEFAULT_MODEL in models
        ollama_error = ""
    except HTTPException as exc:
        models = []
        ollama_model_available = False
        ollama_error = str(exc.detail)
    return {
        "status": "ok",
        "backend": current_backend(),
        "default_backend": current_backend(),
        "openrouter_configured": bool(OPENROUTER_API_KEYS),
        "openrouter_model": OPENROUTER_MODEL,
        "openrouter_key_count": len(OPENROUTER_API_KEYS),
        "openrouter_max_concurrency": OPENROUTER_MAX_CONCURRENCY,
        "openrouter_pool_available": _openrouter_pool.qsize(),
        "ollama_urls": OLLAMA_URLS,
        "ollama_pool_size": len(OLLAMA_URLS),
        "ollama_pool_available": _ollama_pool.qsize(),
        "ollama_model": DEFAULT_MODEL,
        "ollama_model_available": ollama_model_available,
        "ollama_error": ollama_error,
        "synology_chat_configured": synology_chat_configured(),
        "guidance_path": str(GUIDANCE_PATH),
        "timeout_seconds": int(REQUEST_TIMEOUT),
    }


@app.get("/api/usage")
def usage_summary() -> dict[str, Any]:
    with closing(sqlite3.connect(DB_PATH)) as conn:
        conn.row_factory = sqlite3.Row
        cost_rows = conn.execute("SELECT cost_usd FROM cost_log").fetchall()
        generation_rows = conn.execute(
            "SELECT cost_usd FROM generations WHERE cost_usd IS NOT NULL"
        ).fetchall()

    total_cost = round(sum(float(row["cost_usd"]) for row in cost_rows), 6)
    generation_count = len(generation_rows)
    credits: dict[str, Any] | None = None
    if USE_OPENROUTER and OPENROUTER_API_KEYS:
        credits = fetch_openrouter_credits(OPENROUTER_API_KEYS[0])

    return {
        "total_cost_usd": total_cost,
        "generated_count_with_cost": generation_count,
        "credits": credits,
        "backend": current_backend(),
    }


@app.get("/api/generations")
def recent_generations() -> list[dict[str, Any]]:
    return list_recent_generations()


@app.get("/api/generations/jobs/{job_id}")
def generation_job_status(job_id: int) -> dict[str, Any]:
    return get_job(job_id)


@app.get("/api/generations/{generation_id}/markdown")
def download_generation_markdown(generation_id: int) -> PlainTextResponse:
    record = get_generation_record(generation_id)
    markdown = build_generation_markdown(record)
    filename = f"generation-{generation_id}-abstract.md"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return PlainTextResponse(markdown, media_type="text/markdown; charset=utf-8", headers=headers)


@app.post("/api/generations/{generation_id}/notify-writing-club")
def notify_writing_club(generation_id: int) -> dict[str, Any]:
    record = get_generation_record(generation_id)
    result = json.loads(record["result_json"])
    if not result.get("pass_all"):
        raise HTTPException(status_code=400, detail="合格時のみ論文執筆会へ確認依頼を投稿できます。")

    message_text = build_synology_chat_message(record)
    post_synology_chat_message(message_text)
    return {"ok": True, "message": "論文執筆会に確認依頼を投稿しました。"}


@app.post("/api/revise-section")
def revise_section(
    background_tasks: BackgroundTasks,
    field_name: str = Form(...),
    backend: str = Form(""),
    target_language: str = Form("ja"),
    background: str = Form(""),
    problem: str = Form(""),
    objective: str = Form(""),
    challenge: str = Form(""),
    methods: str = Form(""),
    experiment: str = Form(""),
    expected_results: str = Form(""),
    discussion: str = Form(""),
    conclusion: str = Form(""),
    contribution: str = Form(""),
    model: str = Form(DEFAULT_MODEL),
) -> dict[str, Any]:
    if field_name not in SECTION_LABELS:
        raise HTTPException(status_code=400, detail="不正な項目名です。")
    backend_name = resolve_backend_choice(backend)

    request_data = GenerationRequest(
        target_language=target_language.strip() or "ja",
        background=background.strip(),
        problem=problem.strip(),
        objective=objective.strip(),
        challenge=challenge.strip(),
        methods=methods.strip(),
        experiment=experiment.strip(),
        expected_results=expected_results.strip(),
        discussion=discussion.strip(),
        conclusion=conclusion.strip(),
        contribution=contribution.strip(),
        model=model_for_backend(backend_name),
    )

    current_text = getattr(request_data, field_name).strip()
    if len(current_text) < 3:
        raise HTTPException(status_code=400, detail=f"{SECTION_LABELS[field_name]}は 3 文字以上入力してください。")

    guidance_text = load_guidance()
    if backend_name == "openrouter":
        api_key = _openrouter_pool.get()
        try:
            result, generation_id = revise_section_openrouter(request_data, field_name, guidance_text, api_key)
            if generation_id:
                background_tasks.add_task(fetch_and_log_cost, generation_id, api_key, f"section:{field_name}")
            return result
        finally:
            _openrouter_pool.put(api_key)

    ollama_url = _ollama_pool.get()
    try:
        return revise_section_ollama(request_data, field_name, guidance_text, ollama_url)
    finally:
        _ollama_pool.put(ollama_url)


@app.post("/api/evaluate-draft")
def evaluate_draft(
    generation_id: int = Form(0),
    backend: str = Form(""),
    target_language: str = Form("ja"),
    background: str = Form(""),
    problem: str = Form(""),
    objective: str = Form(""),
    challenge: str = Form(""),
    methods: str = Form(""),
    experiment: str = Form(""),
    expected_results: str = Form(""),
    discussion: str = Form(""),
    conclusion: str = Form(""),
    contribution: str = Form(""),
    revised_title: str = Form(""),
    revised_abstract: str = Form(""),
    model: str = Form(DEFAULT_MODEL),
) -> dict[str, Any]:
    backend_name = resolve_backend_choice(backend)
    title_text = revised_title.strip()
    abstract_text = revised_abstract.strip()
    if len(title_text) < 3 or len(abstract_text) < 10:
        raise HTTPException(status_code=400, detail="タイトルは 3 文字以上、アブストラクトは 10 文字以上入力してください。")

    request_data = GenerationRequest(
        target_language=target_language.strip() or "ja",
        background=background.strip(),
        problem=problem.strip(),
        objective=objective.strip(),
        challenge=challenge.strip(),
        methods=methods.strip(),
        experiment=experiment.strip(),
        expected_results=expected_results.strip(),
        discussion=discussion.strip(),
        conclusion=conclusion.strip(),
        contribution=contribution.strip(),
        model=model_for_backend(backend_name),
    )
    result = {
        "revised_title": title_text,
        "revised_abstract": abstract_text,
    }

    if backend_name == "openrouter":
        api_key = _openrouter_pool.get()
        try:
            evaluation = evaluate_generated_outputs_openrouter(request_data, result, api_key)
        finally:
            _openrouter_pool.put(api_key)
    else:
        ollama_url = _ollama_pool.get()
        try:
            evaluation = evaluate_generated_outputs_ollama(request_data, result, ollama_url)
        finally:
            _ollama_pool.put(ollama_url)

    rescored_result = apply_evaluation_result(result, evaluation)
    if generation_id > 0:
        save_rescored_generation_result(generation_id, rescored_result)
    rescored_result["generation_id"] = generation_id if generation_id > 0 else None
    return rescored_result


@app.post("/api/generations")
async def create_generation(
    background_tasks: BackgroundTasks,
    backend: str = Form(""),
    last_name: str = Form(""),
    first_name: str = Form(""),
    target_language: str = Form("ja"),
    background: str = Form(""),
    problem: str = Form(""),
    objective: str = Form(""),
    challenge: str = Form(""),
    methods: str = Form(""),
    experiment: str = Form(""),
    expected_results: str = Form(""),
    discussion: str = Form(""),
    conclusion: str = Form(""),
    contribution: str = Form(""),
    model: str = Form(DEFAULT_MODEL),
) -> dict[str, Any]:
    backend_name = resolve_backend_choice(backend)
    full_name = f"{last_name.strip()}{first_name.strip()}"
    if not full_name:
        raise HTTPException(status_code=400, detail="姓と名を入力してください。")
    payload_values = [
        background.strip(),
        problem.strip(),
        objective.strip(),
        challenge.strip(),
        methods.strip(),
        experiment.strip(),
        expected_results.strip(),
        discussion.strip(),
        conclusion.strip(),
        contribution.strip(),
    ]
    if not any(len(item) >= 3 for item in payload_values):
        raise HTTPException(status_code=400, detail="少なくとも 1 項目は 3 文字以上入力してください。")

    job_title = next((item for item in payload_values if item), "タイトル生成")
    job_id = create_job(job_title[:80], "form")
    background_tasks.add_task(
        process_generation_job,
        job_id,
        backend_name=backend_name,
        name=full_name,
        target_language=target_language.strip() or "ja",
        background=background.strip(),
        problem=problem.strip(),
        objective=objective.strip(),
        challenge=challenge.strip(),
        methods=methods.strip(),
        experiment=experiment.strip(),
        expected_results=expected_results.strip(),
        discussion=discussion.strip(),
        conclusion=conclusion.strip(),
        contribution=contribution.strip(),
        model=model,
    )
    return get_job(job_id)
